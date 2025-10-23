import docx

docpath1 = r"challenge_word_documents/sample1.docx"
docpath2 = r"challenge_word_documents/sample2.docx"
docpath3 = r"challenge_word_documents/sample3.docx"

doc1 = docx.Document(docpath1)
doc2 = docx.Document(docpath2)
doc3 = docx.Document(docpath3)

# Merge the contents of doc2 and doc3 into doc1
for doc in [doc2, doc3]:
	for element in doc.element.body:
		doc1.element.body.append(element)

# Save the merged document
doc1.save("challenge_word_documents/mergedsamples.docx")

